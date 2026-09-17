from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "docs" / "ai_interviewer_presentation.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
SKY = RGBColor(224, 242, 254)
MINT = RGBColor(209, 250, 229)
WHITE = RGBColor(255, 255, 255)
SLATE = RGBColor(71, 85, 105)
LIGHT = RGBColor(241, 245, 249)
ORANGE = RGBColor(249, 115, 22)


def add_text(slide, text, x, y, w, h, size=20, color=NAVY, bold=False,
             align=PP_ALIGN.LEFT, font="Malgun Gothic"):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=21, color=NAVY):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Malgun Gothic"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(13)
        p.bullet = True
    return box


def add_title(slide, title, subtitle=None, number=None):
    add_text(slide, title, Inches(0.7), Inches(0.36), Inches(11.6), Inches(0.6),
             size=28, bold=True)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.05), Inches(1.15), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    if subtitle:
        add_text(slide, subtitle, Inches(0.7), Inches(1.15), Inches(11.4), Inches(0.38),
                 size=13, color=SLATE)
    if number is not None:
        add_text(slide, f"{number:02d}", Inches(12.1), Inches(0.35), Inches(0.5), Inches(0.4),
                 size=14, color=SLATE, align=PP_ALIGN.RIGHT)


def add_footer(slide):
    add_text(slide, "QLoRA 기반 AI 면접관", Inches(0.7), Inches(7.08), Inches(3.0), Inches(0.2),
             size=9, color=SLATE)


def add_card(slide, title, body, x, y, w, h, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(226, 232, 240)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.09), h)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_text(slide, title, x + Inches(0.27), y + Inches(0.2), w - Inches(0.42), Inches(0.35),
             size=18, bold=True)
    add_text(slide, body, x + Inches(0.27), y + Inches(0.65), w - Inches(0.45), h - Inches(0.8),
             size=14, color=SLATE)


def add_flow_box(slide, text, x, y, w, h, fill=SKY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = BLUE
    add_text(slide, text, x + Inches(0.1), y + Inches(0.08), w - Inches(0.2), h - Inches(0.16),
             size=16, bold=True, align=PP_ALIGN.CENTER)


def add_arrow(slide, x, y):
    add_text(slide, "→", x, y, Inches(0.38), Inches(0.38), size=24, color=BLUE,
             bold=True, align=PP_ALIGN.CENTER)


def make_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1. Title
    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    add_text(slide, "QLoRA 기반\nAI 면접관", Inches(0.85), Inches(1.2), Inches(7.0), Inches(1.75),
             size=38, color=WHITE, bold=True)
    add_text(slide, "지원자의 직무와 경험을 바탕으로\n심층 기술면접 질문 한 개를 생성합니다.",
             Inches(0.9), Inches(3.35), Inches(6.3), Inches(0.8), size=20, color=SKY)
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.35), Inches(1.55), Inches(3.8), Inches(2.9))
    pill.fill.solid()
    pill.fill.fore_color.rgb = BLUE
    pill.line.fill.background()
    add_text(slide, "8GB VRAM\nQwen3-1.7B\nQLoRA", Inches(8.7), Inches(2.05), Inches(3.1), Inches(1.8),
             size=24, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "AI Interviewer Project", Inches(0.9), Inches(6.65), Inches(3.0), Inches(0.25),
             size=11, color=SKY)

    # 2. Problem
    slide = prs.slides.add_slide(blank)
    add_title(slide, "문제정의: 왜 AI 면접관인가?", number=2)
    add_card(slide, "채용 환경의 변화", "AI 기반 역량검사와\nAI 면접 활용이 확대", Inches(0.75), Inches(1.8), Inches(3.65), Inches(3.55), BLUE)
    add_card(slide, "취업준비생의 문제", "AI 영상면접과 직무 면접을\n개인 경험에 맞춰 준비하기 어려움", Inches(4.85), Inches(1.8), Inches(3.65), Inches(3.55), ORANGE)
    add_card(slide, "프로젝트 목표", "직무·경험 입력 →\n실제 역량을 검증하는 질문 1개", Inches(8.95), Inches(1.8), Inches(3.65), Inches(3.55), RGBColor(16, 185, 129))
    add_text(slide, "고용노동부·한국고용정보원 2025년 기업 채용동향조사: AI 도구를 채용에 사용하는 기업 중 69.8%가 AI 기반 역량검사를 활용",
             Inches(0.78), Inches(5.8), Inches(11.7), Inches(0.45), size=12, color=SLATE)
    add_footer(slide)

    # 3. Requirement
    slide = prs.slides.add_slide(blank)
    add_title(slide, "요구사항과 제약조건", number=3)
    add_card(slide, "요구사항", "• 직무 + 경험 입력\n• 후보 질문 3개 중 Top-1 선택\n• 최종 출력은 질문 한 개\n• 원본 JSONL은 보존", Inches(0.75), Inches(1.65), Inches(5.6), Inches(3.85), BLUE)
    add_card(slide, "제약조건", "• GPU VRAM 8GB\n• 외부 API·수작업 라벨링 없음\n• 서빙 시 추론은 한 번\n• Full Fine-tuning 지양", Inches(6.95), Inches(1.65), Inches(5.6), Inches(3.85), ORANGE)
    add_text(slide, "해결 방향: 로컬 임베딩으로 자동 Top-1 선택 + 4-bit QLoRA SFT", Inches(1.15), Inches(6.0), Inches(11.0), Inches(0.45),
             size=20, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 4. Top1
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Top-1 우선순위 기준", "현재 기준은 ‘질문 품질’이 아니라 ‘지원자 경험과의 관련성’입니다.", 4)
    add_text(slide, "가장 경험과 관련된 질문", Inches(0.8), Inches(1.75), Inches(4.2), Inches(0.45), size=24, bold=True)
    add_bullets(slide, ["이력에 실제로 적힌 기술과 연결", "프로젝트의 선택·실험·문제 해결을 질문", "이력에 없는 기술을 가정하지 않음"],
                Inches(0.9), Inches(2.3), Inches(5.4), Inches(2.3), size=19)
    add_card(slide, "예시", "경험: Qdrant, BGE 임베딩, Top-k=5\n\n후보: ‘Top-k=5를 어떤 평가 지표로 결정했나요?’\n\n→ 경험의 실제 의사결정을 직접 묻는 후보를 우선", Inches(6.65), Inches(1.8), Inches(5.8), Inches(3.55), RGBColor(16, 185, 129))
    add_footer(slide)

    # 5. Selection
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Top-1 추출 방법", "로컬 BGE 임베딩 기반 자동 pseudo-labeling", 5)
    xs = [0.65, 3.25, 5.85, 8.45]
    labels = ["직무 + 경험", "후보 질문 3개", "BGE 임베딩\nCPU 실행", "최고 유사도\n질문 1개"]
    fills = [SKY, SKY, MINT, RGBColor(254, 242, 242)]
    for x, label, fill in zip(xs, labels, fills):
        add_flow_box(slide, label, Inches(x), Inches(2.65), Inches(2.0), Inches(1.1), fill)
    for x in [2.7, 5.3, 7.9]:
        add_arrow(slide, Inches(x), Inches(3.0))
    add_text(slide, "정규화된 벡터의 내적 = cosine similarity", Inches(3.1), Inches(4.55), Inches(7.2), Inches(0.4),
             size=20, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "모델: BAAI/bge-m3  |  외부 API 없음  |  GPU VRAM 추가 사용 없음", Inches(2.15), Inches(5.45), Inches(9.0), Inches(0.3),
             size=14, color=SLATE, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 6. Data flow
    slide = prs.slides.add_slide(blank)
    add_title(slide, "학습 데이터 흐름", number=6)
    xs = [0.55, 3.15, 5.75, 8.35, 10.95]
    labels = ["원본 JSONL\n후보 3개", "후보 분리", "Top-1 선택", "prompt /\ncompletion 생성", "QLoRA\nSFT 학습"]
    for x, label in zip(xs, labels):
        add_flow_box(slide, label, Inches(x), Inches(2.65), Inches(1.85), Inches(1.1), SKY if x < 8 else MINT)
    for x in [2.55, 5.15, 7.75, 10.35]:
        add_arrow(slide, Inches(x), Inches(3.0))
    add_card(slide, "원본 파일 유지", "datas/ai_interview_sft.jsonl을 수정하지 않고, 노트북 메모리에서만 Top-1 학습 데이터셋을 생성", Inches(2.1), Inches(4.65), Inches(9.1), Inches(1.15), BLUE)
    add_footer(slide)

    # 7. QLoRA
    slide = prs.slides.add_slide(blank)
    add_title(slide, "QLoRA 설계: 8GB 환경", number=7)
    add_card(slide, "기본 모델", "Qwen / Qwen3-1.7B\n\n기본 모델 가중치는 고정", Inches(0.75), Inches(1.65), Inches(3.65), Inches(3.85), BLUE)
    add_card(slide, "양자화", "4-bit NF4\ndouble quantization\n\n메모리 사용량 감소", Inches(4.85), Inches(1.65), Inches(3.65), Inches(3.85), RGBColor(16, 185, 129))
    add_card(slide, "학습", "LoRA rank 16\nbatch size 1\ngradient accumulation 8\nmax length 512", Inches(8.95), Inches(1.65), Inches(3.65), Inches(3.85), ORANGE)
    add_text(slide, "gradient checkpointing + paged_adamw_8bit + completion_only_loss", Inches(1.2), Inches(6.0), Inches(11.0), Inches(0.45),
             size=19, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 8. LoRA comparison
    slide = prs.slides.add_slide(blank)
    add_title(slide, "LoRA와 QLoRA 비교 계획", number=8)
    rows = [
        ("기본 모델", "bf16 / fp16", "4-bit NF4"),
        ("학습 대상", "LoRA adapter", "LoRA adapter"),
        ("VRAM", "상대적으로 큼", "상대적으로 작음"),
        ("8GB 적합성", "OOM 가능성", "실제 학습 권장"),
    ]
    x0, y0 = Inches(1.0), Inches(1.65)
    widths = [Inches(2.7), Inches(4.2), Inches(4.2)]
    headers = ["구분", "LoRA", "QLoRA"]
    for col, (header, width) in enumerate(zip(headers, widths)):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0 + sum(widths[:col]), y0, width, Inches(0.55))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY
        shape.line.color.rgb = WHITE
        add_text(slide, header, x0 + sum(widths[:col]), y0, width, Inches(0.55), size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    for row_i, row in enumerate(rows):
        y = y0 + Inches(0.55 * (row_i + 1))
        for col, (value, width) in enumerate(zip(row, widths)):
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0 + sum(widths[:col]), y, width, Inches(0.55))
            shape.fill.solid()
            shape.fill.fore_color.rgb = LIGHT if row_i % 2 == 0 else WHITE
            shape.line.color.rgb = RGBColor(203, 213, 225)
            add_text(slide, value, x0 + sum(widths[:col]), y, width, Inches(0.55), size=15,
                     bold=(col == 0), align=PP_ALIGN.CENTER)
    add_text(slide, "공정 비교 조건: 동일 JSONL · 동일 Top-1 · 동일 모델 · 동일 rank · 동일 학습 step", Inches(1.25), Inches(5.3), Inches(10.8), Inches(0.4),
             size=18, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 9. Evaluation
    slide = prs.slides.add_slide(blank)
    add_title(slide, "테스트 및 검증", number=9)
    add_card(slide, "기능 검증", "• 후보 3개가 정상 분리되는가?\n• completion이 질문 한 개인가?\n• 번호·해설 없이 출력되는가?", Inches(0.75), Inches(1.65), Inches(3.65), Inches(3.8), BLUE)
    add_card(slide, "자원 비교", "• 최대 VRAM\n• 학습 시간\n• OOM 발생 여부\n• 최종 train loss", Inches(4.85), Inches(1.65), Inches(3.65), Inches(3.8), ORANGE)
    add_card(slide, "질문 품질", "• 경험 관련성\n• 구체성\n• 검증 가능성\n• Before / After 비교", Inches(8.95), Inches(1.65), Inches(3.65), Inches(3.8), RGBColor(16, 185, 129))
    add_footer(slide)

    # 10. Improvements
    slide = prs.slides.add_slide(blank)
    add_title(slide, "개선 방향", number=10)
    add_bullets(slide, [
        "현재 31개 데이터를 수백 건 이상의 직무·경험 사례로 확장",
        "경험 관련성 외에 질문의 구체성과 검증 가능성을 점수화",
        "LoRA와 QLoRA를 동일 조건에서 학습하고 VRAM·시간·품질 비교",
        "질문 한 문장만 출력하도록 후처리 검증 추가",
    ], Inches(1.1), Inches(1.65), Inches(10.8), Inches(3.3), size=22)
    add_text(slide, "결론: 8GB 환경에서는 QLoRA가 실용적인 학습 기준선이며,\n성능 향상의 핵심은 모델 확장보다 데이터 품질과 Top-1 기준 개선입니다.",
             Inches(1.05), Inches(5.35), Inches(11.0), Inches(0.75), size=20, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PATH)
    print(f"Created: {OUTPUT_PATH}")


if __name__ == "__main__":
    make_presentation()
